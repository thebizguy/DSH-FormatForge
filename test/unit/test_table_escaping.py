"""
FF-M-table/audit 回归测试：单元格内容不得撕开 Markdown 表格几何。

背景：``|`` 会伪造新的列边界、换行会伪造新的行——下游 Markdown 渲染器会把
单元格内容当成表格结构（内容欺骗），列数/行数与真实表格不符。
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

import pytest

from core.table_semantics import escape_md_cell, render_markdown_table

pytest.importorskip("openpyxl")


def _cells_of(line: str) -> list[str]:
    """按未转义的 | 切分 Markdown 行（跳过 `\\|`）。"""
    parts = []
    buf = ""
    i = 0
    while i < len(line):
        ch = line[i]
        if ch == "\\" and i + 1 < len(line) and line[i + 1] == "|":
            buf += "|"
            i += 2
            continue
        if ch == "|":
            parts.append(buf)
            buf = ""
            i += 1
            continue
        buf += ch
        i += 1
    parts.append(buf)
    return parts


class TestEscapeMdCell:
    def test_pipe_is_escaped(self):
        assert escape_md_cell("a|b") == "a\\|b"

    def test_newlines_become_br(self):
        assert escape_md_cell("a\nb") == "a<br>b"
        assert escape_md_cell("a\r\nb") == "a<br>b"
        assert escape_md_cell("a\rb") == "a<br>b"

    def test_already_escaped_pipe_not_double_escaped(self):
        assert escape_md_cell("a\\|b") == "a\\|b"

    def test_pipe_after_an_even_backslash_run_is_escaped(self):
        # Two backslashes escape each other in Markdown, so the following pipe
        # is still structural and needs one more backslash.
        assert escape_md_cell("a\\\\|b") == "a\\\\\\|b"

    def test_none_is_empty(self):
        assert escape_md_cell(None) == ""


class TestRenderMarkdownTableGeometry:
    def test_pipe_in_cell_does_not_add_columns(self):
        grid = [["h1", "h2"], ["a|b", "c"]]
        md = render_markdown_table(grid)
        lines = md.split("\n")
        widths = {len(_cells_of(line)) for line in lines}
        assert widths == {4}, f"列边界被撕开: {lines}"

    def test_newline_in_cell_does_not_add_rows(self):
        grid = [["h1", "h2"], ["line1\nline2", "c"]]
        md = render_markdown_table(grid)
        assert len(md.split("\n")) == 3, f"行边界被撕开: {md!r}"
        assert "<br>" in md

    def test_header_pipe_also_escaped(self):
        grid = [["h|1", "h2"], ["a", "b"]]
        lines = render_markdown_table(grid).split("\n")
        assert {len(_cells_of(line)) for line in lines} == {4}


class TestDocxTableCells:
    def test_docx_cell_with_pipe_and_newline(self, tmp_path):
        docx = pytest.importorskip("docx")
        from parsers.docx_parser import DOCXParser

        path = tmp_path / "tbl.docx"
        doc = docx.Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "a|b"
        table.cell(0, 1).text = "第一行"
        table.cell(0, 1).add_paragraph("第二行")
        doc.save(path)

        page = DOCXParser().parse(path)[0]
        table_elems = [e for e in page.elements if e.elementType == "table"]
        assert table_elems, page.elements
        content = table_elems[0].content
        assert "a\\|b" in content
        assert "<br>" in content
        # 单元格换行不得产生额外的表格行
        assert len(content.split("\n")) == 1, content


class TestXlsxTableCells:
    def test_xlsx_cell_with_pipe_and_newline(self, tmp_path):
        import openpyxl

        from parsers.xlsx_parser import XLSXParser

        path = tmp_path / "cells.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["head1", "head2"])
        ws.append(["a|b", "line1\nline2"])
        wb.save(path)

        page = XLSXParser().parse(path)[0]
        table_elems = [e for e in page.elements if e.elementType == "table"]
        assert table_elems, page.elements
        lines = table_elems[0].content.split("\n")
        md_lines = [line for line in lines if line.startswith("|")]
        assert len(md_lines) == 3, lines  # 表头 + 分隔行 + 1 数据行
        assert "a\\|b" in table_elems[0].content
        assert "<br>" in table_elems[0].content
        for line in md_lines:
            assert len(_cells_of(line)) == 4, f"列边界被撕开: {line}"
