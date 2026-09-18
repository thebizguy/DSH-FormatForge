"""
FF-M-docx/audit 回归测试：DOCX 不得静默丢正文。

背景：解析循环只认 body 的直接子节点 ``w:p``/``w:tbl``——
- ``w:sdt``（内容控件/结构化文档标签）里的正文被整段丢弃；
- ``Paragraph.text`` 只拼接直接 ``w:r``，``w:ins`` 追踪插入与超链接文字从正文消失
  （只活在 revisions 元数据里）；
- 单个畸形元素抛异常会废掉整篇文档。
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

import pytest

pytest.importorskip("docx")

from docx import Document  # noqa: E402
from lxml import etree  # noqa: E402

from parsers.docx_parser import DOCXParser  # noqa: E402

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _parse(path):
    pages = DOCXParser().parse(path)
    return pages[0]


def _all_text(page) -> str:
    return "\n".join(e.content for e in page.elements)


class TestStructuredDocumentTags:
    """w:sdt / w:sdtContent 里的正文必须被解析。"""

    def test_sdt_paragraph_is_included(self, tmp_path):
        path = tmp_path / "sdt.docx"
        doc = Document()
        doc.add_paragraph("正文第一段")
        sdt = etree.fromstring(
            f'<w:sdt xmlns:w="{W_NS}"><w:sdtPr><w:tag w:val="t"/></w:sdtPr>'
            "<w:sdtContent><w:p><w:r><w:t>控件里的正文</w:t></w:r></w:p></w:sdtContent></w:sdt>"
        )
        doc.element.body.insert(1, sdt)
        doc.add_paragraph("正文第三段")
        doc.save(path)

        text = _all_text(_parse(path))
        assert "控件里的正文" in text, f"sdt 内容丢失: {text!r}"
        assert "正文第一段" in text and "正文第三段" in text

    def test_sdt_table_is_included(self, tmp_path):
        path = tmp_path / "sdt_tbl.docx"
        doc = Document()
        doc.add_paragraph("前言")
        sdt = etree.fromstring(
            f'<w:sdt xmlns:w="{W_NS}"><w:sdtContent><w:tbl>'
            '<w:tblGrid><w:gridCol w:w="1000"/></w:tblGrid>'
            "<w:tr><w:tc><w:p><w:r><w:t>控件表格单元</w:t></w:r></w:p></w:tc></w:tr>"
            "</w:tbl></w:sdtContent></w:sdt>"
        )
        doc.element.body.insert(1, sdt)
        doc.save(path)

        page = _parse(path)
        assert "控件表格单元" in _all_text(page)
        assert page.hasTable is True

    def test_nested_sdt_is_included(self, tmp_path):
        path = tmp_path / "nested_sdt.docx"
        doc = Document()
        sdt = etree.fromstring(
            f'<w:sdt xmlns:w="{W_NS}"><w:sdtContent><w:sdt><w:sdtContent>'
            "<w:p><w:r><w:t>嵌套控件正文</w:t></w:r></w:p>"
            "</w:sdtContent></w:sdt></w:sdtContent></w:sdt>"
        )
        doc.element.body.insert(0, sdt)
        doc.save(path)

        assert "嵌套控件正文" in _all_text(_parse(path))


class TestTrackedInsertionsInBody:
    """w:ins 插入文本必须出现在正文（当前只出现在 revisions 元数据）。"""

    def test_ins_text_merged_into_body_with_marker(self, tmp_path):
        path = tmp_path / "ins.docx"
        doc = Document()
        para = doc.add_paragraph("原文")
        para._p.append(
            etree.fromstring(
                f'<w:ins xmlns:w="{W_NS}" w:id="1" w:author="Alice" w:date="2026-08-28T10:00:00Z">'
                "<w:r><w:t>插入的内容</w:t></w:r></w:ins>"
            )
        )
        doc.save(path)

        page = _parse(path)
        text = _all_text(page)
        assert "插入的内容" in text, f"插入文本未进正文: {text!r}"
        assert "原文" in text
        # 合并展示但显式标记，便于下游区分未接受的修订
        assert page.elements[0].metadata.get("tracked_insert") is True

    def test_deleted_text_is_not_in_body(self, tmp_path):
        path = tmp_path / "del.docx"
        doc = Document()
        para = doc.add_paragraph("保留文字")
        para._p.append(
            etree.fromstring(
                f'<w:del xmlns:w="{W_NS}" w:id="2" w:author="Bob" w:date="2026-08-28T10:01:00Z">'
                "<w:r><w:delText>被删除的文字</w:delText></w:r></w:del>"
            )
        )
        doc.save(path)

        text = _all_text(_parse(path))
        assert "保留文字" in text
        assert "被删除的文字" not in text

    def test_hyperlink_text_included(self, tmp_path):
        """Paragraph.text 同样漏掉 w:hyperlink 里的文字。"""
        path = tmp_path / "link.docx"
        doc = Document()
        para = doc.add_paragraph("前缀 ")
        para._p.append(
            etree.fromstring(
                f'<w:hyperlink xmlns:w="{W_NS}" r:id="rId1" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                "<w:r><w:t>链接文字</w:t></w:r></w:hyperlink>"
            )
        )
        doc.save(path)

        assert "链接文字" in _all_text(_parse(path))


class TestPerElementIsolation:
    """单个畸形元素不得废掉整篇文档，且必须可见。"""

    def test_malformed_table_does_not_abort_document(self, tmp_path, monkeypatch):
        path = tmp_path / "broken_tbl.docx"
        doc = Document()
        doc.add_paragraph("表格之前")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "单元"
        doc.add_paragraph("表格之后")
        doc.save(path)

        def boom(_table):
            raise ValueError("模拟畸形元素")

        monkeypatch.setattr(DOCXParser, "_extract_table_text", staticmethod(boom))
        page = _parse(path)

        text = _all_text(page)
        assert "表格之前" in text and "表格之后" in text, f"畸形元素废掉了文档: {text!r}"
        assert page.metadata.get("skipped_count") == 1
        assert "tbl" in page.metadata.get("skipped_elements", [])
